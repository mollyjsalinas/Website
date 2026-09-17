"""Parse the BLS OEWS May 2025 all-areas workbook for the six therapy occupations.

Usage: python parse_oews.py <path-to-all_data_M_2025.xlsx>

Reads the 'All May 2025 data' sheet (openpyxl read-only, ~400k rows, ~1-2 min).
Keeps cross-industry rows (NAICS 000000) for:
  national  AREA_TYPE 1                         -> key "US"
  state     AREA_TYPE 2, keyed by PRIM_STATE     -> key "TX"
  metro     AREA_TYPE 4, keyed by AREA=CBSA code -> key "26420"
Fields TOT_EMP, A_MEAN, A_MEDIAN; '*' / '**' / '#' / blank = suppressed -> null.
Output: cache/oews.json  {"period": "May 2025 estimates", "data": {"<key>|<role>": {...}}}
"""
import json, os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(HERE, "cache")
os.makedirs(CACHE, exist_ok=True)

SOC = {
    "29-1123": "pt", "29-1122": "ot", "29-1127": "slp", "29-1181": "aud",
    "31-2021": "pta", "31-2011": "ota",
}

def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip().replace(",", "")
    if s in ("*", "**", "#", "-"):
        return None
    try:
        return int(round(float(s)))
    except ValueError:
        return None

def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    sheet = next(n for n in wb.sheetnames if "data" in n.lower())
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(it)]
    ix = {h: i for i, h in enumerate(hdr)}
    AREA, ATYPE, PSTATE, NAICS, OCC = ix["AREA"], ix["AREA_TYPE"], ix["PRIM_STATE"], ix["NAICS"], ix["OCC_CODE"]
    EMP, MEAN, MED = ix["TOT_EMP"], ix["A_MEAN"], ix["A_MEDIAN"]
    TITLE = ix["AREA_TITLE"]

    data = {}
    titles = {}
    rows = 0
    for row in it:
        rows += 1
        if rows % 100000 == 0:
            print(f"  {rows} rows...", flush=True)
        role = SOC.get(str(row[OCC]))
        if not role or str(row[NAICS]) != "000000":
            continue
        atype = str(row[ATYPE])
        if atype == "1":
            key = "US"
        elif atype == "2":
            key = str(row[PSTATE])
        elif atype == "4":
            key = str(row[AREA]).zfill(5)
        else:
            continue
        titles[key] = row[TITLE]
        data[f"{key}|{role}"] = {
            "employment": num(row[EMP]),
            "annual_mean_wage": num(row[MEAN]),
            "annual_median_wage": num(row[MED]),
        }
    out = {"period": "May 2025 estimates", "source_file": os.path.basename(xlsx),
           "titles": titles, "data": data, "rows_scanned": rows}
    path = os.path.join(CACHE, "oews.json")
    json.dump(out, open(path, "w", encoding="utf-8"), indent=1)
    n_states = len({k.split("|")[0] for k in data if len(k.split("|")[0]) == 2 and k[:2] != "US"})
    n_metros = len({k.split("|")[0] for k in data if len(k.split("|")[0]) == 5})
    print(f"done -> {path}: {len(data)} keys, {n_states} states, {n_metros} metros", flush=True)

if __name__ == "__main__":
    main(sys.argv[1])
